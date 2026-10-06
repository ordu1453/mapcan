"""Экспорт карты в Excel (openpyxl)."""

from __future__ import annotations

import re
from dataclasses import dataclass, field, fields
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from ruamel.yaml import YAML

from ..docmodel import build_view
from ..i18n import _
from ..model import MapCan

CONTENT_COLUMNS = ["device", "frame", "id", "byte", "bits", "signal", "description", "encoding", "values"]


@dataclass
class XlsxStyle:
    """Оформление. Можно переопределить YAML-файлом (``--style``)."""

    font: str = "Calibri"
    font_size: int = 10
    header_fill: str = "1F4E79"
    header_font_color: str = "FFFFFF"
    device_fill: str = "DDEBF7"
    frame_fill: str = "F2F2F2"
    border_color: str = "A6A6A6"
    unused_font_color: str = "A6A6A6"
    columns: list[str] = field(default_factory=lambda: list(CONTENT_COLUMNS))
    max_width: dict[str, int] = field(default_factory=lambda: {
        "device": 22, "frame": 24, "id": 14, "byte": 7, "bits": 8, "signal": 26,
        "description": 40, "encoding": 20, "values": 70,
    })
    bitmap_cell_width: int = 13

    @classmethod
    def load(cls, path: str | Path | None) -> XlsxStyle:
        st = cls()
        if path:
            data = YAML(typ="safe").load(Path(path).read_text(encoding="utf-8")) or {}
            names = {f.name for f in fields(cls)}
            for k, v in data.items():
                if k not in names:
                    raise ValueError(_("неизвестный параметр стиля: {key}", key=k))
                if k == "max_width":
                    st.max_width.update(v)
                else:
                    setattr(st, k, v)
        return st


def _rgb(color: str) -> str:
    return color.lstrip("#").upper()


def _sheet_title(name: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", name)[:31]
    base, i = name, 2
    while name in used:
        suffix = f" ({i})"
        name = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(name)
    return name


class _Writer:
    def __init__(self, cmap: MapCan, style: XlsxStyle, lang: str | None):
        self.view = build_view(cmap, lang)
        self.L = self.view["labels"]
        self.st = style
        side = Side(style="thin", color=_rgb(style.border_color))
        self.border = Border(left=side, right=side, top=side, bottom=side)
        self.font = Font(name=style.font, size=style.font_size)
        self.bold = Font(name=style.font, size=style.font_size, bold=True)
        self.mono = Font(name="Consolas", size=style.font_size)
        self.head_font = Font(name=style.font, size=style.font_size, bold=True,
                              color=_rgb(style.header_font_color))
        self.head_fill = PatternFill("solid", fgColor=_rgb(style.header_fill))
        self.wrap = Alignment(wrap_text=True, vertical="top")
        self.center = Alignment(wrap_text=True, vertical="center", horizontal="center")
        self.wb = Workbook()
        self.titles: set[str] = set()

    # ------------------------------------------------------------------ #

    def run(self, out: Path) -> None:
        self.wb.remove(self.wb.active)
        self.summary_sheet()
        for bus in self.view["buses"]:
            self.addressing_sheet(bus)
            self.exchange_sheet(bus)
            self.content_sheet(bus)
            self.bitmap_sheet(bus)
        self.history_sheet()
        out.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(out)

    def sheet(self, title: str) -> Worksheet:
        return self.wb.create_sheet(_sheet_title(title, self.titles))

    def header(self, ws: Worksheet, row: int, labels: list[str]) -> None:
        for i, lab in enumerate(labels, 1):
            c = ws.cell(row, i, lab)
            c.font, c.fill, c.alignment, c.border = self.head_font, self.head_fill, self.center, self.border
        ws.row_dimensions[row].height = 20

    def put(self, ws, row, col, value, font=None, fill=None, align=None):
        c = ws.cell(row, col, value)
        c.font = font or self.font
        c.alignment = align or self.wrap
        c.border = self.border
        if fill:
            c.fill = PatternFill("solid", fgColor=_rgb(fill))
        return c

    def merge(self, ws, r0, r1, col):
        if r1 > r0:
            ws.merge_cells(start_row=r0, start_column=col, end_row=r1, end_column=col)
            for r in range(r0 + 1, r1 + 1):
                ws.cell(r, col).border = self.border

    def autowidth(self, ws, widths: dict[int, int]) -> None:
        for col, w in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = w

    @staticmethod
    def measure(values, lo=6, hi=60) -> int:
        n = 0
        for v in values:
            for line in str(v or "").split("\n"):
                n = max(n, len(line))
        return max(lo, min(hi, n + 2))

    # ------------------------------------------------------------------ #

    @staticmethod
    def add_logo(ws, path: str, anchor: str, height_px: int = 48) -> None:
        """Логотип (PNG/JPG/GIF) в лист; SVG Excel не поддерживает — пропускаем."""
        if not path or Path(path).suffix.lower() not in (".png", ".jpg", ".jpeg", ".gif"):
            return
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(path)
        except (ImportError, OSError, ValueError):
            return
        if img.height:
            img.width = round(img.width * height_px / img.height)
            img.height = height_px
        ws.add_image(img, anchor)

    def summary_sheet(self) -> None:
        L, d = self.L, self.view["document"]
        ws = self.sheet(L["summary"])
        ws["A1"] = d["title"]
        ws["A1"].font = Font(name=self.st.font, size=16, bold=True)
        self.add_logo(ws, d.get("logo_file", ""), "H1")
        r = 2
        if d["subtitle"]:
            ws.cell(r, 1, d["subtitle"]).font = Font(name=self.st.font, size=12, italic=True)
            r += 1
        r += 1
        for key, val in ((L["version"], d["version"]), (L["date"], d["date"]),
                         (L["author"], d["author"]), ("", "")):
            if val:
                self.put(ws, r, 1, key, self.bold)
                self.put(ws, r, 2, val)
                r += 1
        r += 1
        ws.cell(r, 1, d["bit_note"]).font = Font(name=self.st.font, size=self.st.font_size, italic=True)
        r += 2
        self.header(ws, r, [L["bus"], L["protocol"], L["interface"], L["bitrate"], L["devices"],
                            L["frames_total"].capitalize(), L["signals_total"].capitalize(), L["description"]])
        for bus in self.view["buses"]:
            r += 1
            vals = [bus["title"], bus["protocol"], bus["interface"], bus["bitrate"],
                    bus["stats"]["devices"], bus["stats"]["frames"], bus["stats"]["signals"],
                    bus["description"]]
            for i, v in enumerate(vals, 1):
                self.put(ws, r, i, v)
        self.autowidth(ws, {1: 22, 2: 12, 3: 12, 4: 14, 5: 12, 6: 10, 7: 10, 8: 60})

    # ------------------------------------------------------------------ #

    def addressing_sheet(self, bus: dict) -> None:
        L = self.L
        ws = self.sheet(f"{L['addressing']} {bus['title']}")
        a = bus["addressing"]
        unused = Font(name=self.st.font, size=self.st.font_size, color=_rgb(self.st.unused_font_color))
        if a["kind"] == "matrix":
            heads = [L["message_type"], L["id"], L["purpose"], a["broadcast_label"]] + a["columns"]
            self.header(ws, 1, heads)
            for r, row in enumerate(a["rows"], 2):
                f = self.font if row["used"] else unused
                self.put(ws, r, 1, row["label"], self.bold if row["used"] else unused)
                self.put(ws, r, 2, row["base"], self.mono if row["used"] else unused)
                self.put(ws, r, 3, row["purpose"], f)
                self.put(ws, r, 4, self._cells_text(row["broadcast"]))
                for i, cells in enumerate(row["cells"]):
                    fill = self.st.device_fill if cells else None
                    self.put(ws, r, 5 + i, self._cells_text(cells), fill=fill)
            widths = {1: 12, 2: 9, 3: 18, 4: 22}
            for i in range(len(a["columns"])):
                widths[5 + i] = 24
            self.autowidth(ws, widths)
            ws.freeze_panes = "B2"
        else:
            cols = [("device", L["device"], 16), ("id", L["id"], 13)]
            if a.get("has_pgn"):
                cols.append(("pgn", L["pgn"], 10))
            cols += [("description", L["description"], 50), ("direction", L["direction"], 12),
                     ("period", L["period"], 16)]
            if a.get("has_sender"):
                cols.append(("sender", L["sender"], 20))
            if a.get("has_receivers"):
                cols.append(("receivers", L["receivers"], 30))
            self.header(ws, 1, [c[1] for c in cols])
            for r, row in enumerate(a["rows"], 2):
                for i, (key, _, _) in enumerate(cols, 1):
                    self.put(ws, r, i, row[key], self.mono if key in ("id", "pgn") else None)
            self.autowidth(ws, {i: c[2] for i, c in enumerate(cols, 1)})
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(a['rows']) + 1}"

    def exchange_sheet(self, bus: dict) -> None:
        """Матрица обмена: какое устройство передаёт (TX) и принимает (RX) каждый кадр."""
        L, ex = self.L, bus["exchange"]
        if not ex["columns"] or not ex["rows"]:
            return
        ws = self.sheet(f"{L['exchange']} {bus['title']}")
        self.header(ws, 1, [L["id"], L["frame"], L["description"]] + ex["columns"])
        for r, row in enumerate(ex["rows"], 2):
            self.put(ws, r, 1, row["id"], self.mono)
            self.put(ws, r, 2, row["name"], self.bold)
            self.put(ws, r, 3, row["description"])
            for i, v in enumerate(row["cells"]):
                tx = v == L["tx"]
                self.put(ws, r, 4 + i, v, self.bold if tx else None, self.st.device_fill if tx else None,
                         self.center)
        widths = {1: 13, 2: 24, 3: 50}
        for i in range(len(ex["columns"])):
            widths[4 + i] = 14
        self.autowidth(ws, widths)
        ws.freeze_panes = "D2"

    @staticmethod
    def _cells_text(cells: list[dict]) -> str:
        return "\n".join(f"{c['id']} — {c['description']}" if c["description"] else c["id"] for c in cells)

    # ------------------------------------------------------------------ #

    def content_sheet(self, bus: dict) -> None:
        L, st = self.L, self.st
        ws = self.sheet(f"{L['content']} {bus['title']}")
        cols = [c for c in st.columns if c in CONTENT_COLUMNS]
        heads = {"device": L["device"], "frame": L["frame"], "id": L["id"], "byte": L["byte"],
                 "bits": L["bits"], "signal": L["signal"], "description": L["description"],
                 "encoding": L["encoding"], "values": L["values"]}
        self.header(ws, 1, [heads[c] for c in cols])
        idx = {c: i + 1 for i, c in enumerate(cols)}
        content: dict[str, list] = {c: [] for c in cols}
        row = 2

        groups = []
        if bus["broadcast"]:
            groups.append(({"label": L["broadcast"], "title": "", "description": ""}, bus["broadcast"]))
        for dv in bus["devices"]:
            groups.append((dv, dv["frames"]))

        for dev, frames in groups:
            dev_start = row
            for fr in frames:
                fr_start = row
                if not fr["signals"]:
                    self._cell(ws, row, idx, "description", L["no_signals"], content,
                               font=Font(name=st.font, size=st.font_size, italic=True))
                    row += 1
                for sg in fr["signals"]:
                    sig_start = row
                    lines = sg["values"] or [""]
                    for ln in lines:
                        self._cell(ws, row, idx, "values", ln, content)
                        row += 1
                    sig_end = row - 1
                    desc = sg["description"]
                    if sg["mux"]:
                        desc += f"\n({sg['mux']})"
                    if sg["comment"]:
                        desc += f"\n{L['comment']}: {sg['comment']}"
                    for key, val, font, align in (
                        ("byte", sg["bytes"], self.mono, self.center),
                        ("bits", sg["bits"], self.mono, self.center),
                        ("signal", sg["name"], self.bold, None),
                        ("description", desc.strip(), None, None),
                        ("encoding", "\n".join(sg["encoding"]), None, None),
                    ):
                        self._cell(ws, sig_start, idx, key, val, content, font=font, align=align,
                                   merge_to=sig_end)
                    if "signal" in idx:
                        ws.cell(sig_start, idx["signal"]).fill = PatternFill("solid", fgColor=_rgb(sg["color"]))
                fr_end = row - 1
                ftext = "\n".join(x for x in (
                    fr["title"] or fr["name"],
                    fr["description"],
                    " · ".join(x for x in (fr["type_label"], fr["period"], fr["direction"]) if x),
                    f"DLC {fr['dlc']}",
                    f"{L['pgn']} {fr['pgn']} · {L['priority']} {fr['priority']}" if fr.get("pgn") else "",
                    f"{L['sender']}: {fr['sender']}" if fr.get("sender") else "",
                    f"{L['receivers']}: {fr['receivers']}" if fr.get("receivers") else "",
                ) if x)
                idtext = fr["id"] + (f"\n({fr['id_formula']})" if fr["id_formula"] else "")
                self._cell(ws, fr_start, idx, "frame", ftext, content, fill=st.frame_fill, merge_to=fr_end)
                self._cell(ws, fr_start, idx, "id", idtext, content, font=self.mono, fill=st.frame_fill,
                           merge_to=fr_end, align=Alignment(wrap_text=True, vertical="top", horizontal="center"))
            if row == dev_start:
                continue
            dtext = "\n".join(x for x in (dev.get("label"), dev.get("title"), dev.get("description")) if x)
            self._cell(ws, dev_start, idx, "device", dtext, content, font=self.bold, fill=st.device_fill,
                       merge_to=row - 1)

        ws.freeze_panes = ws.cell(2, idx.get("signal", 1))
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(row - 1, 1)}"
        self.autowidth(ws, {i: self.measure(content[c] + [heads[c]], hi=st.max_width.get(c, 60))
                            for c, i in idx.items()})

    def _cell(self, ws, row, idx, key, val, content, font=None, fill=None, align=None, merge_to=None):
        if key not in idx:
            return
        col = idx[key]
        self.put(ws, row, col, val, font, fill, align)
        content[key].append(val)
        if merge_to is not None:
            self.merge(ws, row, merge_to, col)

    # ------------------------------------------------------------------ #

    def bitmap_sheet(self, bus: dict) -> None:
        L, st = self.L, self.st
        ws = self.sheet(f"{L['bitmaps']} {bus['title']}")
        ws.column_dimensions["A"].width = 8
        for c in range(2, 10):
            ws.column_dimensions[get_column_letter(c)].width = st.bitmap_cell_width
        ws.column_dimensions["J"].width = 50
        thick = Side(style="medium", color="C00000")
        frames = list(bus["broadcast"])
        for dv in bus["devices"]:
            frames += [dict(f, _dev=dv["label"]) for f in dv["frames"]]
        r = 1
        for fr in frames:
            if not fr["grid"]:
                continue
            title = f"{fr['id']}  {fr['name']}"
            if fr.get("_dev"):
                title += f"  ({fr['_dev']})"
            if fr["description"]:
                title += f" — {fr['description']}"
            ws.cell(r, 1, title).font = Font(name=st.font, size=st.font_size + 1, bold=True)
            r += 1
            self.header(ws, r, [L["byte"]] + [str(b) for b in range(7, -1, -1)])
            self.put(ws, r, 10, L["signal"], self.head_font, _rgb(st.header_fill), self.center)
            r += 1
            legend = [f"{s['name']}: {L['byte'].lower()} {s['bytes']}, {L['bits'].lower()} {s['bits']}"
                      for s in fr["signals"]]
            for i, grow in enumerate(fr["grid"]):
                self.put(ws, r, 1, grow["byte"], self.bold, align=self.center)
                for seg in grow["segments"]:
                    c0 = 2 + seg["col"]
                    c = self.put(ws, r, c0, seg["label"] or "", fill=seg["color"], align=self.center,
                                 font=Font(name=st.font, size=st.font_size - 1,
                                           bold=seg["conflict"], color="C00000" if seg["conflict"] else None))
                    if seg["conflict"]:
                        c.border = Border(left=thick, right=thick, top=thick, bottom=thick)
                    if seg["span"] > 1:
                        ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + seg["span"] - 1)
                self.put(ws, r, 10, legend[i] if i < len(legend) else "")
                ws.row_dimensions[r].height = 30
                r += 1
            for extra in legend[len(fr["grid"]):]:
                self.put(ws, r, 10, extra)
                r += 1
            r += 1
        if r == 1:
            ws.cell(1, 1, L["no_signals"])
        ws.cell(r, 1, self.view["document"]["bit_note"]).font = Font(name=st.font, size=st.font_size, italic=True)

    # ------------------------------------------------------------------ #

    def history_sheet(self) -> None:
        L = self.L
        ws = self.sheet(L["changes"])
        self.header(ws, 1, [L["version"], L["date"], L["author"], L["description"]])
        r = 2
        for h in self.view["document"]["history"]:
            changes = h["changes"] or [""]
            for ch in changes:
                self.put(ws, r, 4, ch)
                r += 1
            start = r - len(changes)
            for col, val in ((1, h["version"]), (2, h["date"]), (3, h["author_label"])):
                c = self.put(ws, start, col, val, self.bold if col == 1 else None)
                if col == 3 and h["github_url"]:
                    c.hyperlink = h["github_url"]
                self.merge(ws, start, r - 1, col)
        self.autowidth(ws, {1: 10, 2: 12, 3: 18, 4: 90})
        ws.freeze_panes = "A2"


def export_xlsx(cmap: MapCan, out: str | Path, style: XlsxStyle | None = None, lang: str | None = None) -> Path:
    out = Path(out)
    _Writer(cmap, style or XlsxStyle(), lang).run(out)
    return out
