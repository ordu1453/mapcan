"""Генераторы документов и CLI."""

from __future__ import annotations

import json
import os

import openpyxl
import pytest
from typer.testing import CliRunner

from canmap.cli import app
from canmap.codegen.c import CGen
from canmap.docmodel import build_view
from canmap.export.pdf import export_pdf
from canmap.export.text import render
from canmap.export.xlsx import XlsxStyle, export_xlsx
from canmap.schema import json_schema

from .conftest import DATA, GOLDEN, REFERENCE

UPDATE = os.environ.get("UPDATE_GOLDEN") == "1"


def golden(name: str, actual: str) -> None:
    """Сравнить с эталоном; UPDATE_GOLDEN=1 перезаписывает эталоны."""
    path = GOLDEN / name
    if UPDATE or not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(actual, encoding="utf-8", newline="\n")
    assert actual == path.read_text(encoding="utf-8"), f"{name} отличается от эталона (UPDATE_GOLDEN=1)"


# --------------------------------------------------------------------------- view


def test_view_grid_and_flags(ref_map):
    v = build_view(ref_map)
    bus = v["buses"][0]
    imd = next(d for d in bus["devices"] if d["name"] == "imd")
    state = imd["frames"][0]
    assert state["id"] == "0x196" and state["id_formula"] == "0x180 + node_id"
    byte1 = state["grid"][1]["segments"]
    labels = [s["label"] for s in byte1]
    assert labels == ["internal_error", "self_test_error", "anomaly_error", "timeout_error",
                      "low_bus_voltage_error"]
    parse = next(f for f in bus["devices"][1]["frames"] if f["name"] == "cc_parse_data")
    seg = parse["grid"][1]["segments"][0]
    assert seg["label"] == "voltage_data | temperature_data" and not seg["conflict"]
    assert any("при parse_type = 0" in s["mux"] for s in parse["signals"])


def test_view_hex_values(ref_map):
    v = build_view(ref_map)
    sdo = next(f for d in v["buses"][0]["devices"] for f in d["frames"] if f["name"] == "imd_sdo_request")
    cmd = sdo["signals"][0]
    assert cmd["values"] == ["0x2B – WRITE_2_BYTES"]


def test_view_english(ref_map):
    v = build_view(ref_map, "en")
    assert v["labels"]["signal"] == "Signal"


# --------------------------------------------------------------------------- xlsx


def test_xlsx(ref_map, tmp_path):
    out = export_xlsx(ref_map, tmp_path / "m.xlsx")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == [
        "Сводка", "Адресация Шина 1", "Содержание Шина 1", "Битовые карты Шина 1",
        "Адресация Шина 2", "Содержание Шина 2", "Битовые карты Шина 2", "Изменения",
    ]
    ws = wb["Содержание Шина 1"]
    assert ws.freeze_panes == "F2"
    assert ws.auto_filter.ref.startswith("A1:I")
    assert len(ws.merged_cells.ranges) > 50
    texts = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert "insulation_status" in texts
    assert any(isinstance(t, str) and t.startswith("3 – ALARM") for t in texts)
    addr = wb["Адресация Шина 1"]
    row = [c.value for c in addr[6]]  # PDO TX 1
    assert row[0] == "PDO TX 1" and "0x196" in row[4]


def test_xlsx_custom_columns(ref_map, tmp_path):
    style_file = tmp_path / "s.yaml"
    style_file.write_text("columns: [frame, id, signal, values]\nheader_fill: '00FF00'\n", encoding="utf-8")
    out = export_xlsx(ref_map, tmp_path / "m.xlsx", XlsxStyle.load(style_file))
    ws = openpyxl.load_workbook(out)["Содержание Шина 1"]
    assert [c.value for c in ws[1]] == ["Кадр", "ID", "Сигнал", "Значения"]
    assert ws["A1"].fill.fgColor.rgb.endswith("00FF00")


def test_xlsx_style_unknown_key(tmp_path):
    f = tmp_path / "s.yaml"
    f.write_text("colour: red\n", encoding="utf-8")
    with pytest.raises(ValueError):
        XlsxStyle.load(f)


# --------------------------------------------------------------------------- pdf / md / html


def test_pdf(ref_map, tmp_path):
    out = export_pdf(ref_map, tmp_path / "m.pdf")
    data = out.read_bytes()
    assert data.startswith(b"%PDF") and len(data) > 20_000
    assert data.count(b"/Type /Page") > 5 or data.count(b"/Type/Page") > 5


def test_pdf_custom_template(ref_map, tmp_path):
    tpl = tmp_path / "t.typ"
    tpl.write_text('#let data = json(bytes(sys.inputs.data))\n= #data.document.title\n', encoding="utf-8")
    out = export_pdf(ref_map, tmp_path / "m.pdf", template=tpl)
    assert out.read_bytes().startswith(b"%PDF")


def test_markdown_golden(ref_map):
    golden("canmap.md", render(ref_map, "md"))


def test_html(ref_map):
    html = render(ref_map, "html")
    assert html.startswith("<!doctype html>")
    assert 'id="frame-cc_parse_data"' in html
    assert "&lt;" not in html.split("<style>")[0]  # заголовок без двойного экранирования
    assert "<script>alert" not in html


def test_html_escapes(ref_map):
    m = ref_map.model_copy(deep=True)
    m.document.title = "<script>alert(1)</script>"
    assert "<script>alert(1)" not in render(m, "html")


def test_c_header_golden(ref_map):
    golden("bus1.h", CGen(ref_map, ref_map.buses[0]).header())


# --------------------------------------------------------------------------- schema


def test_schema_accepts_reference():
    jsonschema = pytest.importorskip("jsonschema")
    from canmap.diagnostics import Report
    from canmap.loader import SourceMap, _make_yaml, _plain

    raw = _make_yaml().load(REFERENCE.read_text(encoding="utf-8"))
    data = _plain(raw, (), SourceMap(None), Report())

    def stringify(o):
        if isinstance(o, dict):
            return {str(k): stringify(v) for k, v in o.items()}
        if isinstance(o, list):
            return [stringify(v) for v in o]
        if hasattr(o, "isoformat"):
            return o.isoformat()
        return o

    jsonschema.validate(stringify(data), json_schema())


def test_schema_file_up_to_date():
    from .conftest import ROOT

    on_disk = json.loads((ROOT / "schema" / "canmap.schema.json").read_text(encoding="utf-8"))
    assert on_disk == json_schema(), "обновите: canmap schema -o schema/canmap.schema.json"


# --------------------------------------------------------------------------- CLI

runner = CliRunner()


def test_cli_validate_ok():
    r = runner.invoke(app, ["validate", str(REFERENCE)])
    assert r.exit_code == 0, r.output
    assert "OK" in r.output


def test_cli_validate_errors_json():
    r = runner.invoke(app, ["validate", str(DATA / "bad_map.yaml"), "--format", "json"])
    assert r.exit_code == 1
    items = json.loads(r.output)
    assert any(i["code"] == "number-comma" and i["line"] for i in items)


def test_cli_strict(tmp_path):
    f = tmp_path / "w.yaml"
    f.write_text(REFERENCE.read_text(encoding="utf-8").replace("0: NO_DATA\n", "0: No_Data\n", 1),
                 encoding="utf-8")
    assert runner.invoke(app, ["validate", str(f)]).exit_code == 0
    assert runner.invoke(app, ["validate", str(f), "--strict"]).exit_code == 1


def test_cli_export_refuses_on_errors(tmp_path):
    r = runner.invoke(app, ["export", "md", str(DATA / "bad_map.yaml"), "-o", str(tmp_path / "x.md")])
    assert r.exit_code == 1
    assert not (tmp_path / "x.md").exists()


def test_cli_build(tmp_path):
    r = runner.invoke(app, ["build", str(REFERENCE), "-o", str(tmp_path), "--formats", "md,html,dbc,c,py,xlsx"])
    assert r.exit_code == 0, r.output
    stem = REFERENCE.stem
    for p in (f"{stem}.md", f"{stem}.html", f"{stem}.xlsx", "dbc/bus1.dbc", "c/bus1.h", "c/bus1.c",
              "py/bus1_canmap.py"):
        assert (tmp_path / p).exists(), p


def test_cli_template_and_rules(tmp_path):
    r = runner.invoke(app, ["template", "pdf", "-o", str(tmp_path / "t.typ")])
    assert r.exit_code == 0 and (tmp_path / "t.typ").exists()
    r = runner.invoke(app, ["rules"])
    assert "bit-overlap" in r.output
